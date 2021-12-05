import openpyxl
from openpyxl.chart import LineChart,Reference
from openpyxl import load_workbook
wb=load_workbook(r'./scores.xlsx')
ws=wb['普通推免']
chart=LineChart()
data=Reference(ws,min_row=2,max_row=500,min_col=6,max_col=8)
chart.add_data(data,titles_from_data=True,from_rows=False)
ws.add_chart(chart,'A571')
cate=Reference(ws,min_row=1,max_row=1,min_col=6,max_col=8)
chart.set_categories(cate)
chart.x_axis.title='各类别'
chart.y_axis.title='成绩值'
chart.style=47
wb.save(r'./scores.xlsx')